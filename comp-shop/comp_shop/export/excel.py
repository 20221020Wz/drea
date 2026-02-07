"""Excel导出模块 - 生成标准化的竞品分析Excel"""

import logging
from collections import defaultdict
from pathlib import Path
from typing import Any, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# 样式定义
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
SUBHEADER_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
SUBHEADER_FONT = Font(name="Calibri", bold=True, size=10)
GAP_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
OPPORTUNITY_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


class ExcelExporter:
    """Excel导出器

    生成包含以下工作表的Excel文件：
    1. 各零售商原始数据（按Tab分开）
    2. 横向对比矩阵
    3. Gap分析结果
    4. 统计摘要
    """

    def __init__(self, output_path: str | Path):
        self.output_path = Path(output_path)
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        self.wb = Workbook()
        # 删除默认sheet
        self.wb.remove(self.wb.active)

    def _apply_header_style(self, ws, row: int, max_col: int):
        """应用表头样式"""
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = THIN_BORDER

    def _apply_data_style(self, ws, start_row: int, end_row: int, max_col: int):
        """应用数据区域样式"""
        for row in range(start_row, end_row + 1):
            for col in range(1, max_col + 1):
                cell = ws.cell(row=row, column=col)
                cell.border = THIN_BORDER
                cell.alignment = Alignment(vertical="center", wrap_text=True)

    def _auto_column_width(self, ws, min_width: int = 10, max_width: int = 50):
        """自动调整列宽"""
        for col_idx, col_cells in enumerate(ws.columns, 1):
            max_length = 0
            for cell in col_cells:
                if cell.value:
                    cell_len = len(str(cell.value))
                    max_length = max(max_length, cell_len)
            adjusted = min(max(max_length + 2, min_width), max_width)
            ws.column_dimensions[get_column_letter(col_idx)].width = adjusted

    # ──────────────────── 原始数据表 ────────────────────

    def add_raw_data_sheets(self, products_by_retailer: dict[str, list[dict]]):
        """为每个零售商创建原始数据工作表

        Args:
            products_by_retailer: {retailer_name: [product_dict, ...]}
        """
        retailer_labels = {
            "homedepot": "Home Depot",
            "walmart": "Walmart",
            "lowes": "Lowes",
            "harborfreight": "Harbor Freight",
        }

        columns = [
            ("SKU", "sku"),
            ("Product Name", "name"),
            ("Brand", "brand"),
            ("Price ($)", "price"),
            ("Rating", "rating"),
            ("Reviews", "review_count"),
            ("Type", "type"),
            ("Size", "size"),
            ("Material", "material"),
            ("Handle Type", "handle_type"),
            ("Features", "features"),
            ("In-Store Only", "in_store_only"),
            ("Product URL", "product_url"),
        ]

        for retailer, products in products_by_retailer.items():
            label = retailer_labels.get(retailer, retailer)
            ws = self.wb.create_sheet(title=label[:31])  # Excel sheet name limit

            # 写入表头
            for col_idx, (header, _) in enumerate(columns, 1):
                ws.cell(row=1, column=col_idx, value=header)
            self._apply_header_style(ws, 1, len(columns))

            # 写入数据
            for row_idx, product in enumerate(products, 2):
                specs = product.get("specs", {})
                for col_idx, (_, field) in enumerate(columns, 1):
                    value = specs.get(field) or product.get(field, "")
                    if isinstance(value, list):
                        value = ", ".join(str(v) for v in value)
                    elif isinstance(value, bool):
                        value = "Yes" if value else "No"
                    ws.cell(row=row_idx, column=col_idx, value=value)

            self._apply_data_style(ws, 2, len(products) + 1, len(columns))
            self._auto_column_width(ws)
            ws.auto_filter.ref = ws.dimensions

    # ──────────────────── 横向对比矩阵 ────────────────────

    def add_comparison_matrix(
        self,
        products: list[dict],
        primary_dimension: str,
        secondary_dimension: Optional[str] = None,
        primary_label: str = "Primary",
        secondary_label: str = "Secondary",
    ):
        """创建横向对比矩阵工作表

        将竖向数据转为横向对比格式，行为维度值，列为零售商。

        Args:
            products: 所有产品数据
            primary_dimension: 主维度字段名
            secondary_dimension: 次维度字段名（可选）
            primary_label: 主维度显示名
            secondary_label: 次维度显示名
        """
        ws = self.wb.create_sheet(title="Comparison Matrix")

        # 收集所有零售商和维度值
        retailers = sorted(set(p.get("retailer", "") for p in products))
        dimension_values: dict[str, dict[str, list[dict]]] = defaultdict(
            lambda: defaultdict(list)
        )

        for p in products:
            specs = p.get("specs", {})
            primary_val = str(specs.get(primary_dimension) or p.get(primary_dimension, "Other"))
            if primary_val in ("None", "null", ""):
                primary_val = "Other"
            dimension_values[primary_val][p.get("retailer", "")].append(p)

        # 写入表头
        row = 1
        ws.cell(row=row, column=1, value=primary_label)
        for col_idx, retailer in enumerate(retailers):
            ws.cell(row=row, column=col_idx + 2, value=retailer.upper())
        self._apply_header_style(ws, row, len(retailers) + 1)

        # 写入矩阵数据
        row = 2
        for dim_value in sorted(dimension_values.keys()):
            ws.cell(row=row, column=1, value=dim_value)
            ws.cell(row=row, column=1).font = Font(bold=True)

            for col_idx, retailer in enumerate(retailers):
                products_in_cell = dimension_values[dim_value].get(retailer, [])
                if products_in_cell:
                    # 显示产品数量和代表性产品
                    cell_lines = [f"({len(products_in_cell)} products)"]
                    for p in products_in_cell[:3]:
                        name = p.get("name", "")[:40]
                        price = p.get("price", 0)
                        cell_lines.append(f"• {name} ${price:.2f}")
                    cell_value = "\n".join(cell_lines)
                    ws.cell(row=row, column=col_idx + 2, value=cell_value)
                else:
                    # 标记为Gap
                    ws.cell(row=row, column=col_idx + 2, value="— GAP —")
                    ws.cell(row=row, column=col_idx + 2).fill = GAP_FILL
                    ws.cell(row=row, column=col_idx + 2).font = Font(
                        color="9C0006", bold=True
                    )

            row += 1

        self._apply_data_style(ws, 2, row - 1, len(retailers) + 1)
        self._auto_column_width(ws, min_width=15, max_width=60)

        # 设置行高以适应多行内容
        for r in range(2, row):
            ws.row_dimensions[r].height = 80

    # ──────────────────── Gap分析结果表 ────────────────────

    def add_gap_analysis_sheet(self, gaps: list[dict]):
        """创建Gap分析结果工作表

        Args:
            gaps: Gap机会列表 (GapOpportunity.to_dict() 的结果)
        """
        ws = self.wb.create_sheet(title="Gap Analysis")

        headers = [
            "Priority",
            "Gap Type",
            "Target Retailer",
            "Dimension",
            "Value",
            "Competitors",
            "Competitor Examples",
            "Recommendation",
        ]

        # 表头
        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_idx, value=header)
        self._apply_header_style(ws, 1, len(headers))

        # 数据
        gap_type_labels = {
            "complete_gap": "Complete Gap",
            "brand_gap": "Brand Gap",
            "price_band_gap": "Price Band Gap",
            "feature_gap": "Feature Gap",
        }

        for row_idx, gap in enumerate(gaps, 2):
            ws.cell(row=row_idx, column=1, value=gap.get("priority_score", 0))
            ws.cell(
                row=row_idx,
                column=2,
                value=gap_type_labels.get(gap.get("gap_type", ""), gap.get("gap_type", "")),
            )
            ws.cell(row=row_idx, column=3, value=gap.get("target_retailer", ""))
            ws.cell(row=row_idx, column=4, value=gap.get("dimension", ""))
            ws.cell(row=row_idx, column=5, value=gap.get("dimension_value", ""))
            ws.cell(row=row_idx, column=6, value=gap.get("competitors_count", 0))
            ws.cell(
                row=row_idx,
                column=7,
                value="\n".join(gap.get("competitor_examples", [])),
            )
            ws.cell(row=row_idx, column=8, value=gap.get("recommendation", ""))

            # 按优先级着色
            priority = gap.get("priority_score", 0)
            if priority >= 7:
                ws.cell(row=row_idx, column=1).fill = PatternFill(
                    start_color="FF6B6B", end_color="FF6B6B", fill_type="solid"
                )
            elif priority >= 5:
                ws.cell(row=row_idx, column=1).fill = PatternFill(
                    start_color="FFD93D", end_color="FFD93D", fill_type="solid"
                )

        self._apply_data_style(ws, 2, len(gaps) + 1, len(headers))
        self._auto_column_width(ws, min_width=12, max_width=60)
        ws.auto_filter.ref = ws.dimensions

    # ──────────────────── 统计摘要表 ────────────────────

    def add_summary_sheet(self, summary: dict[str, Any]):
        """创建统计摘要工作表

        Args:
            summary: 来自 GapAnalyzer.get_coverage_summary() 的摘要数据
        """
        ws = self.wb.create_sheet(title="Summary")

        # 零售商概览表
        ws.cell(row=1, column=1, value="Retailer Overview")
        ws.cell(row=1, column=1).font = Font(bold=True, size=14)

        headers = [
            "Retailer",
            "Products",
            "Brands",
            "Price Min",
            "Price Max",
            "Price Avg",
            "Avg Rating",
            "Total Reviews",
        ]
        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=3, column=col_idx, value=header)
        self._apply_header_style(ws, 3, len(headers))

        row = 4
        for retailer, data in summary.items():
            ws.cell(row=row, column=1, value=retailer)
            ws.cell(row=row, column=2, value=data.get("product_count", 0))
            ws.cell(row=row, column=3, value=data.get("brand_count", 0))
            price = data.get("price_range", {})
            ws.cell(row=row, column=4, value=f"${price.get('min', 0):.2f}")
            ws.cell(row=row, column=5, value=f"${price.get('max', 0):.2f}")
            ws.cell(row=row, column=6, value=f"${price.get('avg', 0):.2f}")
            ws.cell(row=row, column=7, value=data.get("avg_rating", 0))
            ws.cell(row=row, column=8, value=data.get("total_reviews", 0))
            row += 1

        self._apply_data_style(ws, 4, row - 1, len(headers))
        self._auto_column_width(ws)

        # 添加柱状图 - 产品数量对比
        chart = BarChart()
        chart.title = "Products by Retailer"
        chart.y_axis.title = "Product Count"
        chart.x_axis.title = "Retailer"
        chart.style = 10

        data_ref = Reference(ws, min_col=2, min_row=3, max_row=row - 1)
        cats_ref = Reference(ws, min_col=1, min_row=4, max_row=row - 1)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        chart.shape = 4

        ws.add_chart(chart, f"A{row + 2}")

    # ──────────────────── 保存 ────────────────────

    def save(self) -> str:
        """保存Excel文件

        Returns:
            保存的文件路径
        """
        self.wb.save(str(self.output_path))
        logger.info(f"Excel文件已保存: {self.output_path}")
        return str(self.output_path)


def export_full_report(
    products: list[dict],
    gaps: list[dict],
    summary: dict,
    primary_dimension: str,
    output_path: str | Path,
    primary_label: str = "Primary Dimension",
) -> str:
    """一键导出完整的竞品分析Excel报告

    Args:
        products: 所有产品数据列表
        gaps: Gap分析结果列表
        summary: 覆盖摘要数据
        primary_dimension: 主分析维度
        output_path: 输出文件路径
        primary_label: 主维度显示名

    Returns:
        保存的文件路径
    """
    exporter = ExcelExporter(output_path)

    # 按零售商分组
    products_by_retailer: dict[str, list[dict]] = defaultdict(list)
    for p in products:
        products_by_retailer[p.get("retailer", "unknown")].append(p)

    # 添加各工作表
    exporter.add_raw_data_sheets(products_by_retailer)
    exporter.add_comparison_matrix(products, primary_dimension, primary_label=primary_label)
    exporter.add_gap_analysis_sheet(gaps)
    exporter.add_summary_sheet(summary)

    return exporter.save()
